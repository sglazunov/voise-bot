"""Раздел «Оговорки: что временно, отключено или скрыто» (протокол CRM 15.09.2026).

Тестировщик разобрал протокол вместе с расшифровкой и показал, что самое
полезное он вытащил из СЫРЫХ реплик, а не из протокола: «инсерты временно
отключил» [20:41], «конфигурации в JSON, БД не трогал — на проде будет БД»
[38:51], «на вкладке импорта есть функционал, скрытый по умолчанию» [30:38],
«третья вкладка — заглушка под Word-журналы» [29:02]. Ни один из этих фактов
не решение, не задача и не ответ на вопрос — и в схеме протокола для них не
было места: демо на 20 минут (2137 слов речи) ужалось в 80 слов темы.
Заодно: «логирование решили не делать» протокол подавал как недоделку
(«отсутствует логирование») — осознанный отказ теперь статус «снято».
"""
from __future__ import annotations

import json

import pytest

from app import analyze, analyze_prompts
from app.analyze import MapNotes, Protocol, _mech_merge, _normalise_caveats


class TestСхемаИПромпт:
    def test_схемы_принимают_caveats(self):
        m = MapNotes.model_validate({"caveats": [{"t": "20:41", "item": "инсерты в импорте",
                                                  "kind": "отключено", "note": "ошибки не готовы"}]})
        assert m.caveats[0].kind == "отключено"
        p = Protocol.model_validate({"caveats": [{"item": "конфиги в JSON", "kind": "прототип"}]})
        assert p.caveats[0].item == "конфиги в JSON"

    def test_промпты_просят_оговорки_на_всех_стадиях(self):
        for tpl in (analyze_prompts._SCHEMA, analyze_prompts._MAP_SCHEMA):
            assert '"caveats"' in tpl and "скрыто" in tpl
        assert "ОГОВОРКИ О СОСТОЯНИИ СИСТЕМЫ" in analyze_prompts._RULES
        assert "caveats" in analyze_prompts._REDUCE_TEMPLATE
        assert "caveats" in analyze_prompts._MERGE_TEMPLATE

    def test_caveats_в_схеме_идут_до_detailed(self):
        """Порядок ключей = порядок генерации: при обрезке ответа теряются
        темы, а не списки. Оговорки — список."""
        s = analyze_prompts._SCHEMA
        assert s.index('"caveats"') < s.index('"detailed"')

    def test_отказ_по_решению_это_снято(self):
        assert "ПО РЕШЕНИЮ" in analyze_prompts._RULES
        assert "«снято»" in analyze_prompts._RULES
        assert "ВМЕСТЕ с ней" in analyze_prompts._RULES     # решение с причиной


class TestНормализация:
    def test_словарный_вид_и_хвост_в_note(self):
        out = _normalise_caveats([
            {"item": "Инсерты в импорте", "kind": "Отключено — пока много ошибок", "note": ""},
            {"item": "Конфигурации в JSON", "kind": "прототип", "note": "на проде будет БД"},
        ])
        assert out[0] == {"item": "Инсерты в импорте", "kind": "отключено",
                          "note": "Пока много ошибок"}
        assert out[1]["kind"] == "прототип"

    def test_пустой_вид_не_теряет_оговорку(self):
        out = _normalise_caveats([{"item": "Скрытая вкладка", "kind": "", "note": ""},
                                  "мусор", {"kind": "риск"}])
        assert out == [{"item": "Скрытая вкладка", "kind": "временно", "note": ""}]

    def test_дубли_после_уплотнения_схлопываются(self):
        out = _normalise_caveats([{"item": "Инсерты отключены", "kind": "отключено"},
                                  {"item": "инсерты отключены.", "kind": "отключено"}])
        assert len(out) == 1

    def test_механическое_слияние_не_теряет_caveats(self):
        a = {"time_range": "00:00–10:00", "caveats": [{"item": "a", "kind": "скрыто"}]}
        b = {"time_range": "10:00–20:00", "caveats": [{"item": "b", "kind": "риск"}]}
        assert [x["item"] for x in _mech_merge(a, b)["caveats"]] == ["a", "b"]

    def test_протокол_несёт_caveats_наружу(self, monkeypatch):
        monkeypatch.setattr(analyze, "MIN_SPEECH_WORDS", 0)
        monkeypatch.setattr(analyze, "_MAX_TOPIC_REGENS", 0)
        proto = {"participants": [], "summary": "s", "detailed": [], "key_thoughts": [],
                 "conclusions": [], "decisions": [], "done_tasks": [], "minor_tasks": [],
                 "tasks": [], "statuses": [],
                 "caveats": [{"item": "Конфигурации интеграций хранятся в JSON-файлах",
                              "kind": "прототип", "note": "БД не трогал; на проде будет БД"}]}

        class B:
            name = "x"; model = "m"
            def complete(self, prompt, max_tokens=2000, force_json=True, **kw):
                return json.dumps(proto, ensure_ascii=False)

        monkeypatch.setattr(analyze.llm, "get_provider_chain", lambda *a, **k: B())
        res = analyze.analyze_transcript(
            "[00:38:51] Константин: сохранение конфигурации я делал не через БД, "
            "каждая интеграция — один JSON")
        assert res["caveats"][0]["kind"] == "прототип"


class TestWord:
    def test_раздел_в_docx(self, tmp_path):
        pytest.importorskip("docx")
        from docx import Document
        from app.docx_export import generate_report
        analysis = {"summary": "s", "detailed": [], "decisions": [], "tasks": [],
                    "minor_tasks": [], "done_tasks": [], "conclusions": [], "key_thoughts": [],
                    "caveats": [{"item": "Инсерты импорта", "kind": "отключено",
                                 "note": "пока не готовы ошибки"},
                                {"item": "Функционал на вкладке импорта", "kind": "скрыто",
                                 "note": ""}]}
        out = tmp_path / "p.docx"
        generate_report(out_path=out, filename="x.mp4", segments=[], analysis=analysis)
        d = Document(str(out))
        heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
        assert "Оговорки: что временно, отключено или скрыто" in heads
        cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
        assert "Инсерты импорта" in cells and "скрыто" in cells
