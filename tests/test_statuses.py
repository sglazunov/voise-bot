"""Раздел «Что спрашивали и что ответили» (по протоколу CRM 01.09.2026).

Разработчик после отпуска прочёл протокол и спросил: «какой документ ищете и
нашли ли?», «что с багом с ролями — есть он или нет?». Ответы в расшифровке
были ([41:31] «документ Константина найти бы… попробую найти»; [34:57]–[35:23]
«тип периода привязан к ролям — ещё не сделано, не сложно, остаётся»), но
протокол их растворил в описании тем. Теперь вопрос → ответ — отдельный раздел.
"""
from __future__ import annotations

import json

import pytest

from app import analyze, analyze_prompts
from app.analyze import MapNotes, Protocol, _mech_merge, _normalise_statuses


class TestСхемаИПромпт:
    def test_схемы_принимают_statuses(self):
        m = MapNotes.model_validate({"statuses": [{"t": "41:31", "item": "документ Константина",
                                                   "status": "не сделано", "note": "Зоя поищет"}]})
        assert m.statuses[0].item == "документ Константина"
        p = Protocol.model_validate({"statuses": [{"item": "баг с ролями", "status": "остаётся в плане"}]})
        assert p.statuses[0].status == "остаётся в плане"

    def test_промпты_просят_ответ_на_каждый_вопрос(self):
        for tpl in (analyze_prompts._SCHEMA, analyze_prompts._MAP_SCHEMA):
            assert '"statuses"' in tpl and "без ответа" in tpl
        assert "ВОПРОС → ОТВЕТ" in analyze_prompts._RULES
        assert "statuses" in analyze_prompts._REDUCE_TEMPLATE
        assert "statuses" in analyze_prompts._MERGE_TEMPLATE

    def test_statuses_в_схеме_идут_до_detailed(self):
        s = analyze_prompts._SCHEMA
        assert s.index('"statuses"') < s.index('"detailed"')

    def test_правило_про_первое_лицо_и_тему(self):
        assert "от первого лица" in analyze_prompts._RULES
        assert "НЕ выводи ответственного из темы" in analyze_prompts._RULES


class TestНормализация:
    def test_словарная_форма_статуса_и_хвост_в_note(self):
        out = _normalise_statuses([
            {"item": "Документ Константина", "status": "Не сделано — Зоя поищет", "note": ""},
            {"item": "Баг с ролями", "status": "остается в плане", "note": "не сложно, смотрели"},
            {"item": "Поиск", "status": "Сделано", "note": "проверить сегодня"},
        ])
        assert out[0] == {"item": "Документ Константина", "status": "не сделано", "note": "Зоя поищет"}
        assert out[1]["status"] == "остаётся в плане" and out[1]["note"] == "не сложно, смотрели"
        assert out[2]["status"] == "сделано"

    def test_пустой_статус_значит_без_ответа(self):
        out = _normalise_statuses([{"item": "Изучалось ли?", "status": "", "note": ""}, "мусор", {"status": "x"}])
        assert out == [{"item": "Изучалось ли?", "status": "без ответа", "note": ""}]

    def test_механическое_слияние_не_теряет_statuses(self):
        a = {"time_range": "00:00–10:00", "statuses": [{"item": "a", "status": "сделано"}]}
        b = {"time_range": "10:00–20:00", "statuses": [{"item": "b", "status": "без ответа"}]}
        assert [x["item"] for x in _mech_merge(a, b)["statuses"]] == ["a", "b"]

    def test_протокол_несёт_statuses_наружу(self, monkeypatch):
        monkeypatch.setattr(analyze, "MIN_SPEECH_WORDS", 0)
        monkeypatch.setattr(analyze, "_MAX_TOPIC_REGENS", 0)
        proto = {"participants": [], "summary": "s", "detailed": [], "key_thoughts": [],
                 "conclusions": [], "decisions": [], "done_tasks": [], "minor_tasks": [], "tasks": [],
                 "statuses": [{"item": "Документ Константина по модулю интеграции",
                               "status": "не сделано", "note": "Зоя попробует найти и положить в Яндекс"}]}

        class B:
            name = "x"; model = "m"
            def complete(self, prompt, max_tokens=2000, force_json=True, **kw):
                return json.dumps(proto)

        monkeypatch.setattr(analyze.llm, "get_provider_chain", lambda *a, **k: B())
        res = analyze.analyze_transcript("[41:31] Зоя Р: документ Константина найти бы, попробую найти")
        assert res["statuses"][0]["status"] == "не сделано"


class TestWord:
    def test_раздел_в_docx(self, tmp_path):
        pytest.importorskip("docx")
        from docx import Document
        from app.docx_export import generate_report
        analysis = {"summary": "s", "detailed": [], "decisions": [], "tasks": [], "minor_tasks": [],
                    "done_tasks": [], "conclusions": [], "key_thoughts": [],
                    "statuses": [{"item": "Документ Константина", "status": "не сделано", "note": "Зоя поищет"},
                                 {"item": "Изучалось ли про роли?", "status": "без ответа", "note": ""}]}
        out = tmp_path / "p.docx"
        generate_report(out_path=out, filename="x.mp4", segments=[], analysis=analysis)
        d = Document(str(out))
        heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
        assert "Что спрашивали и что ответили" in heads
        cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
        assert "Документ Константина" in cells and "без ответа" in cells and "Зоя поищет" in cells
