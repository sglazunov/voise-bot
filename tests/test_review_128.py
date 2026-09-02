"""Ревью 128 протоколов (07.07–02.09): что протоколы врали и как это закрыто.

Пункты ревью: галлюцинация разделов; прошедшее время как задача и
проигнорированное «Нет, не надо»; ответственный не читается из обращения;
ложные срабатывания проверки; шум в задачах; таймкоды [013:01]; имена
«Зояр»; задачи прошлой встречи не сверяются со «Сделано».
"""
from __future__ import annotations

import json

import pytest

from app import analyze, analyze_prompts, llm, meeting_series, names
from app.analyze import (_LABEL_RE, _drop_noise_tasks, _find_support, _norm_t,
                         _refine_tasks, check_topics)

SPEECH = "\n".join([
    "[04:40] Зоя Р: давайте по серверу.",
    "[04:46] Сергей Глазунов: надо обновить сервер и закрыть задачу по конструктору.",
    "[04:57] Зоя Р: да, закрывай.",
    "[07:10] Зоя Р: Паша, можешь собрать список шаблонов писем?",
    "[07:20] Павел Мельников: да, соберу.",
    "[09:01] Мельников Алексей: я хочу, чтобы ты, Паша, поставил задачи по статусам.",
    "[12:30] Кирилл Бубнов: я добавил ссылку на базу знаний в мсу.",
    "[12:35] Зоя Р: нет, не надо в мсу.",
    "[14:00] Кирилл Бубнов: скриншоты по тренажёрам я уже скинул в чат.",
    "[15:00] Зоя Р: у нас тридцать тренажёров в работе, Кирилл ведёт аналитику.",
] + [f"[{20 + i:02d}:00] Зоя Р: обсуждаем реплика номер {i} про тренажёры и аналитику" for i in range(40)])


class TestТаймкодыИМетки:
    def test_ноль_часов_и_чмс_приводятся_к_ммсс(self):
        assert _norm_t("[013:01]") == "13:01"
        assert _norm_t("021:22") == "21:22"
        assert _norm_t("00:40:39") == "40:39"
        assert _norm_t("1:05:03") == "65:03"
        assert _norm_t("12:07") == "12:07"
        assert _norm_t("") is None and _norm_t(None) is None

    def test_метка_с_трёхзначными_минутами_вырезается(self):
        assert _LABEL_RE.sub("", "[013:01] Зоя Р: текст") == "текст"


class TestШумВЗадачах:
    def test_шаблонные_фразы_убираются_и_видны_в_dropped(self):
        res = {"tasks": [{"task": "Написать Виктору или руководителю для получения доступа"},
                         {"task": "Обновить сервер"}],
               "minor_tasks": [{"task": "Добавить в контекст Бота слово «утка»", "owner": "Ани Rie"},
                               {"task": "Включить микрофон перед встречей"},
                               {"task": "Переименовать теги"}],
               "done_tasks": []}
        _drop_noise_tasks(res)
        assert [t["task"] for t in res["tasks"]] == ["Обновить сервер"]
        assert [t["task"] for t in res["minor_tasks"]] == ["Переименовать теги"]
        assert len(res["_dropped"]) == 3

    def test_свой_стоп_лист_из_env(self, monkeypatch):
        monkeypatch.setenv("VTX_TASK_STOPLIST", r"кофе;;\bпиццу\b")
        res = {"tasks": [{"task": "Заказать пиццу"}, {"task": "Сварить кофе"}, {"task": "Сделать отчёт"}],
               "minor_tasks": [], "done_tasks": []}
        _drop_noise_tasks(res)
        assert [t["task"] for t in res["tasks"]] == ["Сделать отчёт"]


class TestОпораРазделов:
    def test_раздел_из_контекста_удаляется_а_речевой_остаётся(self):
        res = {"detailed": [
            {"topic": "Технические требования и структура команды",
             "details": "В команде 8 программистов и 3 веб-дизайнера, выпускают 30–50 тренажёров "
                        "в месяц, планируется импорт из Asana и Trello, интеграция с Jira и Notion."},
            {"topic": "Аналитика тренажёров",
             "details": "Обсуждали тренажёры и аналитику: Кирилл ведёт аналитику, в работе тридцать "
                        "тренажёров, обсуждали реплики по тренажёрам."},
        ]}
        check_topics(res, SPEECH)
        assert [t["topic"] for t in res["detailed"]] == ["Аналитика тренажёров"]
        assert res["_dropped_topics"] == ["Технические требования и структура команды"]

    def test_короткая_речь_не_режет(self):
        res = {"detailed": [{"topic": "Т", "details": "8 программистов, Asana, Trello, Jira, Notion, Figma"}]}
        check_topics(res, "[00:01] коротко")
        assert len(res["detailed"]) == 1


class TestДочитываниеПроверки:
    def _proto(self):
        return {"participants": [{"name": "Павел Мельников"}, {"name": "Кирилл Бубнов"},
                                 {"name": "Зоя Р"}, {"name": "Мельников Алексей"}],
                "tasks": [
                    {"task": "Собрать список шаблонов писем", "owner": ""},
                    {"task": "Поставить задачи по статусам", "owner": "Мельников Алексей"},
                    {"task": "Добавить ссылку на базу знаний", "owner": ""},
                    {"task": "Скинуть скриншоты по тренажёрам в чат", "owner": ""},
                ],
                "minor_tasks": [], "done_tasks": [], "decisions": []}

    def _ver(self):
        def ok(q, t):
            return {"ok": True, "quote": q, "t": t, "source": "transcript", "owner_ok": False, "match": "verbatim"}
        return {"tasks": [ok("Паша, можешь собрать список шаблонов писем?", "07:10"),
                          ok("я хочу, чтобы ты, Паша, поставил задачи по статусам", "09:01"),
                          ok("я добавил ссылку на базу знаний в мсу", "12:30"),
                          ok("скриншоты по тренажёрам я уже скинул в чат", "14:00")],
                "minor_tasks": [], "done_tasks": [], "decisions": []}

    def test_ответственный_из_обращения(self):
        res, ver = self._proto(), self._ver()
        _refine_tasks(res, ver, SPEECH)
        by = {t["task"]: t["owner"] for t in res["tasks"]}
        assert by["Собрать список шаблонов писем"] == "Павел Мельников"
        assert by["Поставить задачи по статусам"] == "Павел Мельников"   # не говорящий
        assert ver["tasks"][0]["owner_source"] == "обращение"

    def test_отказ_в_следующей_реплике_снимает_подтверждение(self):
        res, ver = self._proto(), self._ver()
        _refine_tasks(res, ver, SPEECH)
        i = [t["task"] for t in res["tasks"]].index("Добавить ссылку на базу знаний")
        assert ver["tasks"][i]["ok"] is False
        assert "не надо в мсу" in ver["tasks"][i]["note"]

    def test_прошедшее_время_уезжает_в_сделано(self):
        res, ver = self._proto(), self._ver()
        _refine_tasks(res, ver, SPEECH)
        assert [t["task"] for t in res["done_tasks"]] == ["Скинуть скриншоты по тренажёрам в чат"]
        assert ver["done_tasks"][0]["note"] == "по цитате — уже сделано"
        assert len(res["tasks"]) == len(ver["tasks"]) == 3

    def test_опора_ищется_по_основам_когда_модель_цитату_не_нашла(self):
        hit = _find_support("Обновить сервер и закрыть задачу по конструктору", SPEECH)
        assert hit and hit["t"] == "04:46" and "обновить сервер" in hit["quote"]
        assert _find_support("Купить слона в Африке", SPEECH) is None

    def test_verify_protocol_подхватывает_опору_без_модели(self, monkeypatch):
        class B:
            name = "x"; model = "m"
            def complete(self, prompt, max_tokens=2000, force_json=True, **kw):
                return json.dumps({"items": []})     # модель ничего не нашла
        monkeypatch.setattr(analyze.llm, "get_provider_chain", lambda *a, **k: B())
        res = {"tasks": [{"task": "Обновить сервер и закрыть задачу по конструктору", "owner": ""}],
               "minor_tasks": [], "done_tasks": [], "decisions": [], "detailed": []}
        out = analyze.verify_protocol(res, SPEECH)
        v = out["verification"]["tasks"][0]
        assert v["ok"] and v["match"] == "approx" and v["t"] == "04:46"


class TestОбращение:
    KNOWN = ["Павел Мельников", "Наталья П", "Кирилл Бубнов"]

    def test_уменьшительное_сводится_к_участнику(self):
        assert names.addressee("Паша, можешь собрать список", self.KNOWN) == "Павел Мельников"
        assert names.addressee("Это нужно будет, Наташа, кстати, поправить", self.KNOWN) == "Наталья П"
        assert names.addressee("[14:53] Зоя Р: Кирилл, можешь роль там делать?", self.KNOWN) == "Кирилл Бубнов"

    def test_не_имя_и_первое_лицо_не_считаются(self):
        assert names.addressee("Нет, не надо в мсу", self.KNOWN) == ""
        assert names.addressee("Хорошо, посмотрим", self.KNOWN) == ""
        assert names.addressee("Зоя, я сделаю этот список", self.KNOWN + ["Зоя Р"]) == ""

    def test_без_известных_отдаётся_полное_имя(self):
        assert names.addressee("Саша, глянь миграцию", []) == "Александр"

    def test_зояр_сводится_к_зоя_р(self):
        assert names.canonical("Зояр", ["Зоя Р", "Кирилл Бубнов"]) == "Зоя Р"


class TestПереносЗадачПрошлойВстречи:
    def test_статусы_прошлых_задач(self, monkeypatch, tmp_path):
        monkeypatch.setattr(meeting_series, "find", lambda u, t: ("crm", {"last": {
            "date": "25.08.2026", "job_id": "old",
            "tasks": [{"task": "Переименовать теги карточек", "owner": "Кирилл"},
                      {"task": "Создать админа для миграции", "owner": "Саша"},
                      {"task": "Найти документ Константина", "owner": ""}]}}))
        analysis = {"done_tasks": [{"task": "Теги карточек переименованы"}],
                    "tasks": [{"task": "Создать временного админа для миграции"}],
                    "minor_tasks": [], "statuses": []}
        out = meeting_series.carry_over("u", "CRM", analysis, job_id="new")
        st = {i["task"]: i["status"] for i in out["items"]}
        assert st["Переименовать теги карточек"] == "закрыта"
        assert st["Создать админа для миграции"] == "в работе"
        assert st["Найти документ Константина"] == "без упоминания"
        assert out["date"] == "25.08.2026"

    def test_память_этого_же_протокола_не_сверяется(self, monkeypatch):
        monkeypatch.setattr(meeting_series, "find", lambda u, t: ("k", {"last": {
            "job_id": "same", "tasks": [{"task": "x y z"}]}}))
        assert meeting_series.carry_over("u", "T", {"tasks": []}, job_id="same") is None


class TestОжиданиеИПромпт:
    def test_окно_ожидания_лимита_расширено(self):
        assert llm.KEY_WAIT_SEC >= 240 and llm.KEY_TOTAL_WAIT_SEC >= 900

    def test_правила_про_прошедшее_время_и_контекст(self):
        r = analyze_prompts._RULES
        assert "ПРОШЕДШЕЕ ВРЕМЯ — НЕ ЗАДАЧА" in r and "detailed ТОЛЬКО ИЗ РЕЧИ" in r


class TestWordЭкспорт:
    def _doc(self, tmp_path, analysis):
        pytest.importorskip("docx")
        from docx import Document
        from app.docx_export import generate_report
        out = tmp_path / "p.docx"
        base = {"summary": "s", "detailed": [], "decisions": [], "tasks": [], "minor_tasks": [],
                "done_tasks": [], "conclusions": [], "key_thoughts": []}
        generate_report(out_path=out, filename="x.mp4", segments=[], analysis={**base, **analysis})
        d = Document(str(out))
        return d, [p.text for p in d.paragraphs], [c.text for t in d.tables for r in t.rows for c in r.cells]

    def test_черновик_в_заголовке_при_откате(self, tmp_path):
        _d, paras, _cells = self._doc(tmp_path, {"_fallback": ["gemini: 429"]})
        assert any(p.startswith("Протокол встречи — ЧЕРНОВИК") for p in paras)
        assert any(p.startswith("ЧЕРНОВИК:") for p in paras)

    def test_задачи_без_ответственного_отдельным_блоком(self, tmp_path):
        ok = {"ok": True, "quote": "надо обновить сервер", "t": "04:46", "source": "transcript",
              "owner_ok": True, "match": "verbatim"}
        d, paras, cells = self._doc(tmp_path, {
            "tasks": [{"task": "Обновить сервер", "owner": "Сергей"}, {"task": "Собрать шаблоны", "owner": ""}],
            "verification": {"mode": "strict", "tasks": [ok, dict(ok, owner_ok=False)],
                             "minor_tasks": [], "done_tasks": [], "decisions": []}})
        heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
        assert "Задачи без ответственного — назначьте исполнителя" in heads
        assert any("Обновить сервер" in c for c in cells)
        assert not any("Собрать шаблоны" in c for c in cells)
        assert any("Собрать шаблоны" in p for p in paras)

    def test_перенос_и_отказ_видны(self, tmp_path):
        bad = {"ok": False, "quote": "я добавил ссылку", "t": "12:30", "source": "transcript",
               "owner_ok": False, "match": "verbatim", "note": "дальше в разговоре: «нет, не надо в мсу»"}
        d, paras, cells = self._doc(tmp_path, {
            "tasks": [{"task": "Добавить ссылку", "owner": ""}],
            "verification": {"mode": "strict", "tasks": [bad], "minor_tasks": [], "done_tasks": [], "decisions": []},
            "_carried": {"date": "25.08.2026", "items": [{"task": "Найти документ", "owner": "", "status": "без упоминания"}]}})
        assert any("не надо в мсу" in p for p in paras)
        assert "Найти документ" in cells and "без упоминания" in cells


class TestПоПротоколу0209:
    """Сравнение двух протоколов «Встреча лидеров» 02.09 (до/после правок)."""

    def test_усечённое_обращение_ань_лиз_серёж(self):
        known = ["Анна Румянцева", "Елизавета", "Сергей Beck"]
        assert names.addressee("Ань, ты сможешь прописать, разработать текст", known) == "Анна Румянцева"
        assert names.addressee("Лиз, скинь фильм", known) == "Елизавета"
        assert names.addressee("Серёж, глянь курс", known) == "Сергей Beck"
        assert names.addressee("Да, конечно, у меня есть фильм", known) == ""

    def test_решение_дублирующее_задачу_убирается(self):
        res = {"decisions": ["Елизавета отправит скан договора, оригинал — на Открытое шоссе",
                             "Убрать кнопку «Волонтёры» с сайта"],
               "tasks": [], "minor_tasks": [{"task": "Отправить скан договора"}], "done_tasks": []}
        analyze._dedup_decisions(res)
        assert res["decisions"] == ["Убрать кнопку «Волонтёры» с сайта"]
        assert "скан договора" in res["_dropped"][0]

    def test_ожидание_лимита_не_больше_трёх_раундов(self, monkeypatch):
        calls = []

        class Hot:
            name = "gemini"
            def __init__(self, model=None, api_key=None, extra=None): pass
            def complete(self, *a, **k):
                calls.append(1); raise RuntimeError("HTTP 429: try again in 0s")

        monkeypatch.setattr(llm, "_is_rate_limit", lambda e: True)
        monkeypatch.setattr(llm, "_cooldown_from", lambda e, default=0.0: 0.0)
        monkeypatch.setattr(llm.time, "sleep", lambda s: None)
        prov = llm._RotatingProvider(Hot, "m", [("k1", "")])
        with pytest.raises(RuntimeError, match="суточная квота"):
            prov.complete("x")
        assert len(calls) == llm.KEY_MAX_WAITS + 1
